"""Phase 4 regression: convocation registrations + Excel import + record admin.

Covers: fresh and legacy migrations, template download, upload validation and
rejections, preview immutability, commit counts, repeated/conflicting imports,
replace mode, corrections, status transitions with audit, forced mid-import
rollback, a 6,000-row import, search/pagination, unauthorized + CSRF rejection,
and preservation of volunteer/QR/gate/attendance data.

Run:  python scripts/test_exam_registration.py   (from app/)
Disposable temp databases only - never touches app/lgu.db.
"""

import io
import json
import os
import re
import sqlite3
import sys
import tempfile

APP_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, APP_DIR)

TMP = tempfile.mkdtemp(prefix="lgu-p4-")
os.environ["LGU_TESTING"] = "1"
os.environ["LGU_DB_PATH"] = os.path.join(TMP, "fresh.db")

import config  # noqa: E402
config.require_test_db()
import database  # noqa: E402
import exam_auth  # noqa: E402
import exam_import  # noqa: E402
import exam_registrations as regx  # noqa: E402
import migrations  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from web.routes import create_app  # noqa: E402

failures = 0
PW = "correct-horse-battery"
STAFF = "examregofficer"
PRESERVE = ["students", "guests", "gates", "scans", "broadcasts",
            "settings", "convocations"]


def check(label, cond, detail=""):
    global failures
    if cond:
        print("ok   " + label)
    else:
        failures += 1
        print("FAIL " + label + (" - " + str(detail) if detail else ""))


def raw(sql, args=()):
    c = sqlite3.connect(config.DB_PATH)
    c.row_factory = sqlite3.Row
    try:
        return [dict(r) for r in c.execute(sql, args).fetchall()]
    finally:
        c.close()


def run(sql, args=()):
    c = sqlite3.connect(config.DB_PATH)
    try:
        c.execute(sql, args)
        c.commit()
    finally:
        c.close()


def dump(names):
    out = {}
    c = sqlite3.connect(config.DB_PATH)
    for t in names:
        out[t] = c.execute(f"SELECT * FROM {t} ORDER BY rowid").fetchall()
    c.close()
    return out


def csrf_of(html):
    m = re.search(r'name="csrf_token" value="([^"]+)"', html)
    return m.group(1) if m else ""


def xlsx(rows, header=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "registrations"
    ws.append(list(header) if header else list(regx.CANONICAL_COLUMNS))
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(c, token, data, filename="import.xlsx"):
    return c.post("/exam/registrations/import",
                  data={"csrf_token": token,
                        "file": (io.BytesIO(data), filename)},
                  content_type="multipart/form-data")


def batch_id_of(resp):
    m = re.search(r"/import/(\d+)", resp.headers.get("Location", ""))
    return int(m.group(1)) if m else 0


def batch_row(bid):
    return raw("SELECT * FROM import_batches WHERE id = ?", (bid,))[0]


def commit(c, token, bid, mode="insert"):
    return c.post(f"/exam/registrations/import/{bid}/commit",
                  data={"csrf_token": token, "mode": mode},
                  follow_redirects=True)


def html(resp):
    return resp.get_data(as_text=True)


def result_of(bid):
    return json.loads(batch_row(bid)["result_json"])


def count_regs():
    return raw("SELECT COUNT(*) n FROM convocation_registrations")[0]["n"]


# ---------------------------------------------------------------- boot
database.init_db()
database.seed()
app = create_app()
app.config["TESTING"] = True
client = app.test_client()
database.init_db()  # re-run: migrations idempotent
exam_auth.create_staff(STAFF, "Exam Registrar", PW)

# 1. Migrations: fresh database carries the phase 4 tables at latest version.
latest = len(migrations.MIGRATIONS)
check("fresh: user_version is latest",
      raw("PRAGMA user_version")[0]["user_version"] == latest)
tables = {r["name"] for r in raw("SELECT name FROM sqlite_master WHERE type='table'")}
check("fresh: phase 4 tables exist",
      {"convocation_registrations", "import_batches", "import_staging",
       "registration_audit"} <= tables)

# 2. Legacy database at the previous schema version upgrades cleanly.
legacy = os.path.join(TMP, "legacy.db")
c = sqlite3.connect(legacy)
c.executescript(database.SCHEMA)
for v, sql in migrations.MIGRATIONS[:-1]:
    c.executescript(sql)
    c.execute(f"PRAGMA user_version = {v}")
c.commit()
c.close()
config.DB_PATH = legacy
check("legacy: sits at previous version",
      raw("PRAGMA user_version")[0]["user_version"] == latest - 1)
database.init_db()
legacy_version = raw("PRAGMA user_version")[0]["user_version"]
legacy_tables = {r["name"] for r in
                 raw("SELECT name FROM sqlite_master WHERE type='table'")}
check("legacy: migrates to latest version", legacy_version == latest)
check("legacy: phase 4 tables created by migration",
      {"convocation_registrations", "import_batches"} <= legacy_tables)
database.init_db()
check("legacy: migration re-run is a no-op",
      raw("PRAGMA user_version")[0]["user_version"] == latest)
config.DB_PATH = os.path.join(TMP, "fresh.db")

# 3. Snapshot of data this feature must never touch.
PRE = dump(PRESERVE)

# 4. Authorization: nothing reachable without an exam staff session.
fresh = app.test_client()
check("unauth: registrations list redirects",
      fresh.get("/exam/registrations").status_code == 302)
check("unauth: template download redirects",
      fresh.get("/exam/registrations/template").status_code == 302)
check("unauth: upload page redirects",
      fresh.get("/exam/registrations/import").status_code == 302)
check("unauth: upload POST rejected (CSRF guard)",
      upload(fresh, "", b"x").status_code == 403)
check("unauth: commit POST rejected (CSRF guard)",
      fresh.post("/exam/registrations/import/1/commit",
                 data={"csrf_token": "", "mode": "insert"},
                 follow_redirects=False).status_code == 403)
with fresh.session_transaction() as s:
    s["lgu_admin"] = "1"
check("volunteer session cannot open exam portal",
      fresh.get("/exam/registrations").status_code == 302)

# 5. Login, CSRF enforcement with a valid session, template download.
r = client.get("/exam/login")
ltok = csrf_of(html(r))
r = client.post("/exam/login",
                data={"username": STAFF, "password": PW, "csrf_token": ltok})
check("staff login succeeds", r.status_code == 302, r.status_code)
# Login rotates the session token: fetch a fresh one from an authed page.
CSRF = csrf_of(html(client.get("/exam/registrations/import")))
check("fresh csrf token after login", bool(CSRF))
r = upload(client, "", b"x")
check("POST without csrf token rejected", r.status_code == 403, r.status_code)
r = upload(client, "wrong-token", b"x")
check("POST with wrong csrf token rejected", r.status_code == 403, r.status_code)

r = client.get("/exam/registrations/template")
check("template downloads (200, xlsx, attachment)",
      r.status_code == 200 and r.data[:2] == b"PK"
      and "attachment" in r.headers.get("Content-Disposition", ""))
twb = load_workbook(io.BytesIO(r.data))
check("template has canonical header",
      twb.sheetnames[0] == "registrations"
      and [c.value for c in twb["registrations"][1]] == list(regx.CANONICAL_COLUMNS))
notes = " ".join(str(c.value) for row in twb["notes"].iter_rows()
                 for c in row if c.value)
check("template documents assumptions",
      "registration_number" in notes and "student_name" in notes
      and len(notes) > 200, len(notes))

# 6. Happy-path import: preview is read-only, commit inserts with leading zeros.
ROWS5 = [
    ["000001", "a1@uni.edu.pk", "04551", "Ali Noor", "Computer Science"],
    ["000002", "a2@uni.edu.pk", "04552", "Sara Khan", "Physics"],
    ["000003", "a3@uni.edu.pk", "04553", "Omar Farooq", ""],
    ["000004", "a4@uni.edu.pk", "04554", "Hina Aslam", "Mathematics"],
    ["000005", "a5@uni.edu.pk", "04555", "Bilal Ahmed", "Chemistry"],
]
T0 = dump(["convocation_registrations"])
r = upload(client, CSRF, xlsx(ROWS5))
check("upload redirects to preview", r.status_code == 302, r.status_code)
b1 = batch_id_of(r)
check("batch created", b1 > 0, r.headers.get("Location"))
prev = client.get(f"/exam/registrations/import/{b1}")
check("preview renders with confirmable count",
      prev.status_code == 200 and "Confirm import (5 rows)" in html(prev),
      prev.status_code)
check("preview marked preview-only", "Preview only" in html(prev))
check("preview does not mutate registrations",
      dump(["convocation_registrations"]) == T0)
b1r = batch_row(b1)
check("batch counters correct",
      (b1r["total_rows"], b1r["blank_rows"], b1r["valid_rows"],
       b1r["invalid_rows"], b1r["status"]) == (5, 0, 5, 0, "preview"), b1r)
r = commit(client, CSRF, b1, "insert")
check("commit succeeds", r.status_code == 200 and "Import committed" in html(r))
row1 = raw("SELECT * FROM convocation_registrations WHERE registration_number='000001'")
check("leading zeros round-trip",
      row1 and row1[0]["roll_number"] == "04551"
      and row1[0]["registration_number"] == "000001"
      and row1[0]["status"] == "active")
check("5 rows inserted", count_regs() == 5, count_regs())
check("import_insert audited per row",
      raw("SELECT COUNT(*) n FROM registration_audit WHERE action='import_insert'")[0]["n"] == 5)
check("batch committed with result",
      batch_row(b1)["status"] == "committed"
      and result_of(b1)["inserted"] == 5)
check("audit_log import_committed once",
      raw("SELECT COUNT(*) n FROM audit_log WHERE action='import_committed'")[0]["n"] == 1)

# 7. File-level rejections never create a batch.
n_batches = raw("SELECT COUNT(*) n FROM import_batches")[0]["n"]
r = upload(client, CSRF, xlsx([["000009", "x@uni.edu.pk", "09999", "X", ""]],
                              header=["registration_number", "roll_number",
                                      "student_name", "department"]))
check("missing required column rejected",
      r.status_code == 200 and "Missing required column(s): university_email" in html(r))
r = upload(client, CSRF, b"this is not a zip file")
check("non-xlsx bytes rejected",
      r.status_code == 200 and "Not a valid .xlsx" in html(r))
r = upload(client, CSRF, b"")
check("empty upload rejected",
      r.status_code == 200 and "is empty" in html(r))
exam_import.MAX_DATA_ROWS = 3
try:
    r = upload(client, CSRF, xlsx(ROWS5))
finally:
    exam_import.MAX_DATA_ROWS = 6500
check("row limit enforced",
      r.status_code == 200 and "More than 3 data rows" in html(r), html(r)[:300])
exam_import.MAX_UPLOAD_BYTES = 400
try:
    r = upload(client, CSRF, xlsx(ROWS5))
finally:
    exam_import.MAX_UPLOAD_BYTES = 5 * 1024 * 1024
check("size limit enforced",
      r.status_code == 200 and "larger than" in html(r), html(r)[:300])
check("rejected uploads create no batches",
      raw("SELECT COUNT(*) n FROM import_batches")[0]["n"] == n_batches)

# 8. Messy file: blanks, invalid rows, in-file duplicates with row errors.
MESSY = [
    ["000011", "p1@uni.edu.pk", "04561", "Messy One", "CS"],
    [" "] * 5,
    ["000012", "bad-email", "04562", "Messy Two", "CS"],
    ["000013", "p3@uni.edu.pk", "04563", "", "CS"],
    ["000014", "p4@uni.edu.pk", "04564", "Messy Four", "CS"],
    ["000015", "p5@uni.edu.pk", "04561", "Messy Five", "CS"],
    ["R" * 65, "p6@uni.edu.pk", "04565", "Messy Six", "CS"],
    ["000016", "p7@uni.edu.pk", "04566", True, "CS"],
]
T1 = dump(["convocation_registrations"])
r = upload(client, CSRF, xlsx(MESSY))
b2 = batch_id_of(r)
b2r = batch_row(b2)
check("messy counters (8/1/2/4/1/0)",
      (b2r["total_rows"], b2r["blank_rows"], b2r["valid_rows"],
       b2r["invalid_rows"], b2r["duplicate_rows"], b2r["conflict_rows"])
      == (8, 1, 2, 4, 1, 0), b2r)
prev = client.get(f"/exam/registrations/import/{b2}")
check("preview offers only the 2 valid rows",
      "Confirm import (2 rows)" in html(prev))
check("row errors shown on preview",
      "valid email address" in html(prev) and "duplicates row" in html(prev)
      and "student_name is required" in html(prev))
check("messy preview still immutable",
      dump(["convocation_registrations"]) == T1)
commit(client, CSRF, b2, "insert")
res = result_of(b2)
check("messy commit counts",
      (res["inserted"], res["skipped_invalid"], res["skipped_duplicate"])
      == (2, 4, 1), res)
check("7 registrations after messy import", count_regs() == 7, count_regs())

# 8b. department column is optional in the header.
r = upload(client, CSRF,
           xlsx([["000031", "od@uni.edu.pk", "04581", "No Dept", ""]],
                header=["registration_number", "university_email",
                        "roll_number", "student_name"]))
b_od = batch_id_of(r)
check("missing department column accepted",
      b_od > 0 and batch_row(b_od)["valid_rows"] == 1, b_od)
commit(client, CSRF, b_od, "insert")
check("row without department imports empty",
      raw("SELECT department FROM convocation_registrations"
          " WHERE registration_number='000031'")[0]["department"] == "")

# 9. Repeated import of the same file: everything conflicts, nothing changes.
r = upload(client, CSRF, xlsx(MESSY))
b3 = batch_id_of(r)
b3r = batch_row(b3)
check("re-upload marks existing rows as conflicts",
      (b3r["valid_rows"], b3r["conflict_rows"], b3r["invalid_rows"],
       b3r["duplicate_rows"]) == (0, 2, 4, 1), b3r)
prev = client.get(f"/exam/registrations/import/{b3}")
check("conflict-only batch still confirmable",
      "Confirm import (2 rows)" in html(prev))
commit(client, CSRF, b3, "insert")
res = result_of(b3)
check("repeat import leaves rows unchanged",
      res["inserted"] == 0 and res["unchanged"] == 2, res)
check("still 8 registrations", count_regs() == 8, count_regs())

# 10. Cross-record conflict: identifier owned by a different registration.
r = upload(client, CSRF,
           xlsx([["000999", "z@uni.edu.pk", "04561", "Cross One", "CS"]]))
b4 = batch_id_of(r)
check("foreign roll number detected as conflict",
      batch_row(b4)["conflict_rows"] == 1 and batch_row(b4)["valid_rows"] == 0)
commit(client, CSRF, b4, "insert")
res = result_of(b4)
check("conflicting row skipped in insert mode",
      res["skipped_conflict"] == 1 and res["inserted"] == 0, res)
check("conflicting registration not inserted",
      not raw("SELECT id FROM convocation_registrations"
              " WHERE registration_number='000999'"))

# 11. Replace mode updates an existing row; insert mode skips the same row.
REPL = [["000002", "a2@uni.edu.pk", "04552", "Sara Khan Updated", "Mathematics"]]
r = upload(client, CSRF, xlsx(REPL))
b5 = batch_id_of(r)
check("changed existing row staged as conflict",
      batch_row(b5)["conflict_rows"] == 1)
commit(client, CSRF, b5, "replace")
res = result_of(b5)
check("replace mode updates matching row",
      res["updated"] == 1 and res["inserted"] == 0, res)
row2 = raw("SELECT * FROM convocation_registrations"
           " WHERE registration_number='000002'")[0]
check("replace applied new values",
      row2["student_name"] == "Sara Khan Updated"
      and row2["department"] == "Mathematics" and row2["status"] == "active")
check("replace audited with before/after",
      raw("SELECT COUNT(*) n FROM registration_audit"
          " WHERE action='import_update'")[0]["n"] == 1)
r = upload(client, CSRF, xlsx(
    [["000002", "a2@uni.edu.pk", "04552", "Sara Khan Second", "Physics"]]))
b6 = batch_id_of(r)
commit(client, CSRF, b6, "insert")
check("insert mode skips differing existing row",
      result_of(b6)["skipped_conflict"] == 1, result_of(b6))
check("name still updated after insert-mode batch",
      raw("SELECT student_name FROM convocation_registrations"
          " WHERE registration_number='000002'")[0]["student_name"]
      == "Sara Khan Updated")

# 12. Staff correction: audited, validated, registration number immutable.
rid = raw("SELECT id FROM convocation_registrations"
          " WHERE registration_number='000003'")[0]["id"]
r = client.get(f"/exam/registrations/{rid}")
check("detail page renders", r.status_code == 200 and "000003" in html(r))
r = client.post(f"/exam/registrations/{rid}/correct",
                data={"csrf_token": CSRF, "university_email": "a3@uni.edu.pk",
                      "roll_number": "04553", "student_name": "Omar Farooq",
                      "department": "Statistics"}, follow_redirects=True)
upd = raw("SELECT * FROM convocation_registrations WHERE id=?", (rid,))[0]
check("correction saved and shown",
      r.status_code == 200 and upd["department"] == "Statistics"
      and "Saved" in html(r))
last_action = raw("SELECT action FROM registration_audit WHERE registration_id=?"
                  " ORDER BY id DESC", (rid,))[0]["action"]
check("correction audited", last_action == "correct", last_action)
r = client.post(f"/exam/registrations/{rid}/correct",
                data={"csrf_token": CSRF, "university_email": "broken",
                      "roll_number": "04553", "student_name": "Omar Farooq",
                      "department": "Statistics"})
check("invalid correction rejected (400 + message)",
      r.status_code == 400 and "valid email address" in html(r), r.status_code)
check("invalid correction did not persist",
      raw("SELECT university_email FROM convocation_registrations WHERE id=?",
          (rid,))[0]["university_email"] == "a3@uni.edu.pk")
r = client.post(f"/exam/registrations/{rid}/correct",
                data={"csrf_token": CSRF, "registration_number": "HACKED",
                      "university_email": "a3@uni.edu.pk",
                      "roll_number": "04553", "student_name": "Omar Farooq",
                      "department": "Commerce"}, follow_redirects=True)
after = raw("SELECT registration_number, department FROM convocation_registrations"
            " WHERE id=?", (rid,))[0]
check("registration number not editable via correction",
      after["registration_number"] == "000003"
      and after["department"] == "Commerce")

# 13. Status transitions with confirmation + audit; invalid transitions blocked.
rid4 = raw("SELECT id FROM convocation_registrations"
           " WHERE registration_number='000004'")[0]["id"]
r = client.get(f"/exam/registrations/{rid4}/status?action=cancel")
check("cancel confirmation page",
      r.status_code == 200 and "Cancel registration" in html(r))
r = client.post(f"/exam/registrations/{rid4}/status",
                data={"csrf_token": CSRF, "action": "cancel"},
                follow_redirects=True)
check("cancel applied",
      raw("SELECT status FROM convocation_registrations WHERE id=?",
          (rid4,))[0]["status"] == "cancelled")
r = client.post(f"/exam/registrations/{rid4}/status",
                data={"csrf_token": "nope", "action": "reinstate"})
check("status change requires valid csrf", r.status_code == 403, r.status_code)
r = client.post(f"/exam/registrations/{rid4}/status",
                data={"csrf_token": CSRF, "action": "cancel"})
check("invalid transition rejected (400)", r.status_code == 400, r.status_code)
check("status unchanged after invalid transition",
      raw("SELECT status FROM convocation_registrations WHERE id=?",
          (rid4,))[0]["status"] == "cancelled")
client.post(f"/exam/registrations/{rid4}/status",
            data={"csrf_token": CSRF, "action": "reinstate"},
            follow_redirects=True)
check("reinstate returns to active",
      raw("SELECT status FROM convocation_registrations WHERE id=?",
          (rid4,))[0]["status"] == "active")
client.post(f"/exam/registrations/{rid4}/status",
            data={"csrf_token": CSRF, "action": "revoke"}, follow_redirects=True)
check("revoke applied",
      raw("SELECT status FROM convocation_registrations WHERE id=?",
          (rid4,))[0]["status"] == "revoked")
r = client.post(f"/exam/registrations/{rid4}/status",
                data={"csrf_token": CSRF, "action": "reinstate"})
check("revoke is terminal (400)", r.status_code == 400, r.status_code)
check("still revoked after blocked reinstate",
      raw("SELECT status FROM convocation_registrations WHERE id=?",
          (rid4,))[0]["status"] == "revoked")
trail = {t["action"] for t in
         raw("SELECT action FROM registration_audit WHERE registration_id=?",
             (rid4,))}
check("cancel/reinstate/revoke all audited",
      {"cancel", "reinstate", "revoke"} <= trail, trail)
detail_html = html(client.get(f"/exam/registrations/{rid4}"))
check("revoked badge + audit trail rendered",
      "badge bad" in detail_html and "revoke" in detail_html)

# 14. Forced mid-import failure: full rollback, batch marked failed.
NEW3 = [["000020", "f1@uni.edu.pk", "04571", "Fail One", "CS"],
        ["000021", "f2@uni.edu.pk", "04572", "Fail Two", "CS"],
        ["000022", "f3@uni.edu.pk", "04573", "Fail Three", "CS"]]
r = upload(client, CSRF, xlsx(NEW3))
b7 = batch_id_of(r)
check("rollback batch staged", batch_row(b7)["valid_rows"] == 3)
count_before = count_regs()
audit_before = raw("SELECT COUNT(*) n FROM registration_audit")[0]["n"]
orig_insert = regx.insert
calls = {"n": 0}


def boom(db, conv, clean, batch_id, actor):
    calls["n"] += 1
    if calls["n"] >= 2:
        raise RuntimeError("forced mid-import failure")
    return orig_insert(db, conv, clean, batch_id, actor)


regx.insert = boom
raised = False
try:
    exam_import.commit_batch(b7, STAFF, "insert")
except RuntimeError:
    raised = True
finally:
    regx.insert = orig_insert
check("forced failure raised out of commit", raised and calls["n"] == 2)
check("all inserts rolled back", count_regs() == count_before,
      (count_regs(), count_before))
check("audit rows rolled back too",
      raw("SELECT COUNT(*) n FROM registration_audit")[0]["n"] == audit_before)
check("batch marked failed", batch_row(b7)["status"] == "failed")
check("failure audited in audit_log",
      raw("SELECT COUNT(*) n FROM audit_log WHERE action='import_failed'")[0]["n"] >= 1)
again = ""
try:
    exam_import.commit_batch(b7, STAFF, "insert")
except ValueError as exc:
    again = str(exc)
check("failed batch cannot be committed again", "already failed" in again, again)
check("no partial rows persisted",
      not raw("SELECT id FROM convocation_registrations"
              " WHERE registration_number='000020'"))

# 15. Scale: 6,000-row import, then a full re-upload as conflicts.
bulk = [[f"R{i:05d}", f"bulk{i}@uni.edu.pk", f"0{i:05d}",
         f"Bulk Student {i}", ""] for i in range(6000)]
r = upload(client, CSRF, xlsx(bulk))
b8 = batch_id_of(r)
b8r = batch_row(b8)
check("6000 rows staged as valid",
      (b8r["total_rows"], b8r["valid_rows"], b8r["invalid_rows"],
       b8r["blank_rows"]) == (6000, 6000, 0, 0), b8r)
prev = client.get(f"/exam/registrations/import/{b8}")
check("6000-row preview confirmable",
      "Confirm import (6000 rows)" in html(prev))
t_before = count_regs()
commit(client, CSRF, b8, "insert")
check("6000 rows inserted", result_of(b8)["inserted"] == 6000, result_of(b8))
check("count after bulk import", count_regs() == t_before + 6000,
      (count_regs(), t_before))
lead = raw("SELECT roll_number FROM convocation_registrations"
           " WHERE registration_number='R00000'")[0]
check("bulk leading zeros preserved", lead["roll_number"] == "000000", lead)
r = upload(client, CSRF, xlsx(bulk))
b9 = batch_id_of(r)
b9r = batch_row(b9)
check("bulk re-upload conflicts on all 6000",
      b9r["conflict_rows"] == 6000 and b9r["valid_rows"] == 0, b9r)
commit(client, CSRF, b9, "insert")
res = result_of(b9)
check("bulk repeat leaves everything unchanged",
      res["unchanged"] == 6000 and res["inserted"] == 0, res)

# 16. Search, status filter, pagination, history.
page = html(client.get("/exam/registrations?q=04551"))
check("search by roll number",
      "000001" in page and "000002" not in page)
page = html(client.get("/exam/registrations?status=revoked"))
check("status filter isolates revoked",
      "000004" in page and "000001" not in page)
page = html(client.get("/exam/registrations"))
check("pagination present for large lists", "Next" in page)
check("status summary line shown", "Active" in page)
check("404 for unknown registration",
      client.get("/exam/registrations/999999").status_code == 404)
hist = html(client.get("/exam/registrations/history"))
check("history lists committed and failed batches",
      "committed" in hist and "failed" in hist)

# 17. DB-level uniqueness (leading identifiers stay unique per convocation).
conv_id = raw("SELECT id FROM convocations ORDER BY id LIMIT 1")[0]["id"]


def expect_unique(label, sql, args):
    try:
        run(sql, args)
        check(label, False, "insert succeeded")
    except sqlite3.IntegrityError as exc:
        check(label, "UNIQUE" in str(exc).upper(), str(exc))


expect_unique(
    "duplicate registration_number blocked",
    "INSERT INTO convocation_registrations"
    " (convocation_id, registration_number, university_email, roll_number,"
    "  student_name, created_at, updated_at)"
    " VALUES (?, '000001', 'dup1@uni.edu.pk', '09001', 'Dup', 'n', 'n')",
    (conv_id,))
expect_unique(
    "duplicate roll_number blocked",
    "INSERT INTO convocation_registrations"
    " (convocation_id, registration_number, university_email, roll_number,"
    "  student_name, created_at, updated_at)"
    " VALUES (?, '000777', 'dup2@uni.edu.pk', '04551', 'Dup', 'n', 'n')",
    (conv_id,))
expect_unique(
    "duplicate university_email blocked",
    "INSERT INTO convocation_registrations"
    " (convocation_id, registration_number, university_email, roll_number,"
    "  student_name, created_at, updated_at)"
    " VALUES (?, '000778', 'a1@uni.edu.pk', '09002', 'Dup', 'n', 'n')",
    (conv_id,))
check("failed inserts left no rows",
      count_regs() == 6008, count_regs())

# 18. Audit actors and preservation of volunteer/QR/gate/attendance data.
actors = {r["actor"] for r in raw("SELECT DISTINCT actor FROM registration_audit")}
check("all registration audit actors are the exam staff", actors == {STAFF}, actors)
check("volunteer/QR/gate/attendance data untouched",
      dump(PRESERVE) == PRE)
dash = html(client.get("/exam", follow_redirects=True))
check("dashboard links to registrations", "/exam/registrations" in dash)

print("\n%s" % ("PASS" if failures == 0 else "%d FAILURE(S)" % failures))
sys.exit(1 if failures else 0)
