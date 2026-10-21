"""Excel (.xlsx) import for medalist records: preview then commit.

Canonical template columns (no official Examination Department format was
supplied - documented as an assumption in the template's notes sheet):
department, student_name, award_category, position, reference_number.
position and reference_number are optional; extra columns are tolerated and
reported, never trusted. The file is INSERT-ONLY: rows conflicting with an
existing award slot are reported and skipped, and nothing is ever deleted -
corrections happen through the medalist record screens.

Pictures are image assets, not text columns: the spreadsheet never carries
files. Upload a picture on the medalist record screen after committing; rows
imported without pictures are flagged on the dashboard until one is added.

Flow: parse + validate the whole file in memory, then write batch + staging
rows in ONE transaction (preview never touches official records). Commit
re-evaluates every clean row against the current DB inside a single write
transaction; any unexpected error rolls the whole commit back and marks the
batch failed.
"""

import hashlib
import io
import json
import os
import re
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

import database
import exam_auth
import medalists as med

# Medalist lists are per-department award records - far smaller than the
# registration roster. 1,000 rows with headroom inside a 5 MB workbook.
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_DATA_ROWS = 1000
MAX_ERROR_ROWS = 200
SAMPLE_ROWS = 15


class ImportRejected(ValueError):
    """File-level rejection: no batch row is created."""


def _cell_text(value):
    """Canonical text for a cell, or None when the value is malformed."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return None
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, str):
        return value
    return None


def _normalise_header(value):
    text = ("" if value is None else str(value)).strip().lower()
    return re.sub(r"[\s\-]+", "_", text)


def _safe_filename(name):
    name = os.path.basename((name or "").replace("\\", "/")).strip()
    name = re.sub(r"[^\w.\- ]", "_", name)[:120]
    return name or "upload.xlsx"


def _insert_row(db, convocation_id, clean, batch_id, actor):
    """Indirection so tests can force a mid-commit failure."""
    return med.insert(db, convocation_id, clean, actor, batch_id)


def build_template():
    """Standard .xlsx template: header-only medalists sheet + notes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "medalists"
    ws.append(list(med.FIELDS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDE", (24, 28, 18, 16, 22)):
        ws.column_dimensions[col].width = width

    notes = wb.create_sheet("notes")
    for row in (
        ("Medalist import - canonical template", ""),
        ("", ""),
        ("Required columns", "department, student_name, award_category"),
        ("Optional columns", "position (e.g. 1, 2, 3 or '1st'), reference_number"),
        ("Award categories", "gold, silver, bronze or any other officially approved category - the sheet is not limited to three per department."),
        ("Reference number", "Optional. Any approved identifier that matches this convocation's registration_number or roll_number. No official identifier format has been confirmed by LGU - unmatched references are warned about, never rejected. Format cells as text to keep leading zeros."),
        ("Pictures", "Photos are NOT spreadsheet columns. Commit the list first, then upload each approved picture on the medalist record screen - every image is validated server-side."),
        ("Example row (for reference only - delete before uploading)", "Computer Science | Jane Doe | gold | 1 | REG-9001"),
        ("", ""),
        ("Assumption", "No official Examination Department medalist sheet was provided. This template defines the canonical columns; map the issued sheet to these five headers before uploading."),
        ("Scope", "Rows are imported for the currently active convocation only and start as unpublished drafts."),
        ("Safety", "Preview never changes official records. Commit inserts new draft records only; rows that collide with an existing award slot are reported and skipped; nothing is ever deleted."),
    ):
        notes.append(row)
    notes.column_dimensions["A"].width = 46
    notes.column_dimensions["B"].width = 96

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_upload(file_storage, convocation_id, actor):
    """Validate an uploaded .xlsx and stage it as a preview batch.

    Returns (batch_id, summary). Raises ImportRejected for file-level
    problems, in which case no batch row exists.
    """
    data = file_storage.read(MAX_UPLOAD_BYTES + 1)
    if not data:
        raise ImportRejected("The uploaded file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ImportRejected(
            f"File is larger than the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit.")
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ImportRejected(
            "Not a valid .xlsx file. Old .xls files and CSVs are not supported - "
            "please use the downloadable template.")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportRejected("The file could not be read as an Excel workbook.")
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise ImportRejected("The first worksheet is empty.")
        index, ignored = {}, []
        for pos, value in enumerate(header):
            name = _normalise_header(value)
            if name in med.FIELDS and name not in index:
                index[name] = pos
            elif name:
                ignored.append(name)
        missing = [c for c in med.REQUIRED_FIELDS if c not in index]
        if missing:
            raise ImportRejected(
                "Missing required column(s): " + ", ".join(missing) +
                ". Download the standard template for the exact headers.")

        db = database.get_db()
        try:
            db_slots = {(r[0].lower(), r[1].lower(), r[2].lower())
                        for r in db.execute(
                            "SELECT department, award_category, position"
                            " FROM medalists WHERE convocation_id = ?",
                            (convocation_id,))}
        finally:
            db.close()

        staged, seen = [], {"slot": {}, "name": {}}
        blank = total = 0
        for row_no, row in enumerate(rows, start=2):
            total += 1
            if total > MAX_DATA_ROWS:
                raise ImportRejected(
                    f"More than {MAX_DATA_ROWS} data rows. Split the file.")
            cells = row or ()
            if not any(v is not None and str(v).strip() for v in cells):
                blank += 1
                continue
            raw, errors = {}, []
            for name in med.FIELDS:
                pos = index.get(name)
                text = _cell_text(cells[pos]) if pos is not None and pos < len(cells) else ""
                if text is None:
                    errors.append(f"{name}: unsupported cell value.")
                    text = ""
                raw[name] = text
            clean, row_errors = med.clean_row(raw)
            errors.extend(row_errors)
            status = "valid" if not errors else "invalid"
            if not errors:
                slot = (clean["department"].lower(),
                        clean["award_category"].lower(),
                        clean["position"].lower())
                name_key = (clean["department"].lower(),
                            clean["student_name"].lower())
                if slot in seen["slot"]:
                    status = "duplicate"
                    errors.append(
                        "Same award slot as row %d of this file "
                        "(department + category + position)."
                        % seen["slot"][slot])
                elif name_key in seen["name"]:
                    status = "duplicate"
                    errors.append(
                        "Same student as row %d of this file in %s. If these "
                        "are two different people, confirm their reference "
                        "numbers first." % (seen["name"][name_key],
                                            clean["department"]))
                else:
                    seen["slot"][slot] = row_no
                    seen["name"][name_key] = row_no
                    if slot in db_slots:
                        status = "conflict"
                        errors.append(
                            "That award slot already has a record for this "
                            "convocation. Correct the existing record instead.")
            staged.append((row_no, clean, status, errors))

        counts = {"valid": 0, "invalid": 0, "duplicate": 0, "conflict": 0}
        for _, _, status, _ in staged:
            counts[status] += 1
        now = med.utc_now()
        # Fresh connection: batch + staging rows land in one transaction so
        # a failed preview leaves no partial batch behind.
        db = database.get_db()
        try:
            cur = db.execute(
                "INSERT INTO medalist_import_batches"
                " (convocation_id, filename, file_sha256, status, total_rows,"
                "  blank_rows, valid_rows, invalid_rows, duplicate_rows,"
                "  conflict_rows, notes, created_by, created_at)"
                " VALUES (?, ?, ?, 'preview', ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (convocation_id, _safe_filename(file_storage.filename),
                 hashlib.sha256(data).hexdigest(), total, blank,
                 counts["valid"], counts["invalid"], counts["duplicate"],
                 counts["conflict"],
                 "Ignored columns: " + ", ".join(sorted(set(ignored))) if ignored else "",
                 str(actor or "")[:64], now))
            batch_id = cur.lastrowid
            db.executemany(
                "INSERT INTO medalist_import_staging"
                " (batch_id, row_no, department, student_name,"
                "  award_category, position, reference_number, status,"
                "  errors_json, created_at)"
                " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                [(batch_id, row_no, clean["department"], clean["student_name"],
                  clean["award_category"], clean["position"],
                  clean["reference_number"], status,
                  json.dumps(errors, ensure_ascii=False), now)
                 for row_no, clean, status, errors in staged])
            db.commit()
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()
        return batch_id, {"total_rows": total, "blank_rows": blank, **counts}
    finally:
        wb.close()


def commit_batch(batch_id, actor):
    """Apply staged valid rows as draft records in one transaction.

    Returns result counts. Insert-only by design: invalid, duplicate and
    conflicting rows are skipped and reported; nothing is ever deleted or
    overwritten. Any unexpected error rolls the whole commit back.
    """
    db = database.get_db()
    txn_started = False  # only a real transaction failure marks the batch failed
    try:
        batch = db.execute("SELECT * FROM medalist_import_batches WHERE id = ?",
                           (batch_id,)).fetchone()
        if batch is None:
            raise ValueError("Import batch not found.")
        if batch["status"] != "preview":
            raise ValueError(f"Batch is already {batch['status']}.")
        convocation_id = batch["convocation_id"]
        skipped = {s: db.execute(
            "SELECT COUNT(*) FROM medalist_import_staging"
            " WHERE batch_id = ? AND status = ?",
            (batch_id, s)).fetchone()[0]
                   for s in ("invalid", "duplicate", "conflict")}
        rows = db.execute(
            "SELECT * FROM medalist_import_staging WHERE batch_id = ?"
            " AND status = 'valid' ORDER BY row_no",
            (batch_id,)).fetchall()

        counts = {"inserted": 0, "skipped_invalid": skipped["invalid"],
                  "skipped_duplicate": skipped["duplicate"],
                  "skipped_conflict": skipped["conflict"]}
        db.execute("BEGIN IMMEDIATE")
        txn_started = True
        for row in rows:
            clean = {k: row[k] for k in med.FIELDS}
            exists = db.execute(
                "SELECT 1 FROM medalists WHERE convocation_id = ?"
                " AND department = ? COLLATE NOCASE"
                " AND award_category = ? COLLATE NOCASE"
                " AND position = ? COLLATE NOCASE LIMIT 1",
                (convocation_id, clean["department"],
                 clean["award_category"], clean["position"])).fetchone()
            if exists:
                counts["skipped_conflict"] += 1
                continue
            _insert_row(db, convocation_id, clean, batch_id, actor)
            counts["inserted"] += 1

        result = {**counts, "mode": "insert"}
        db.execute(
            "UPDATE medalist_import_batches SET status = 'committed',"
            " result_json = ?, committed_at = ? WHERE id = ?",
            (json.dumps(result), med.utc_now(), batch_id))
        exam_auth.audit(db, actor, "medalist_import_committed",
                        f"batch {batch_id}: {counts['inserted']} inserted, "
                        f"{counts['skipped_invalid']} invalid, "
                        f"{counts['skipped_duplicate']} duplicate, "
                        f"{counts['skipped_conflict']} conflict skipped")
        db.commit()
        return result
    except Exception:
        db.rollback()
        if txn_started:
            try:
                db.execute(
                    "UPDATE medalist_import_batches SET status = 'failed',"
                    " result_json = ? WHERE id = ? AND status = 'preview'",
                    (json.dumps({"error": "commit rolled back"}), batch_id))
                exam_auth.audit(db, actor, "medalist_import_failed",
                                f"batch {batch_id}: commit rolled back, no"
                                " records changed")
                db.commit()
            except Exception:
                db.rollback()
        raise
    finally:
        db.close()


def get_batch(db, batch_id):
    return db.execute("SELECT * FROM medalist_import_batches WHERE id = ?",
                      (batch_id,)).fetchone()


def error_rows(db, batch_id, limit=MAX_ERROR_ROWS):
    rows = db.execute(
        "SELECT * FROM medalist_import_staging WHERE batch_id = ?"
        " AND status IN ('invalid', 'duplicate', 'conflict')"
        " ORDER BY row_no LIMIT ?", (batch_id, limit)).fetchall()
    return [{**dict(r), "errors": json.loads(r["errors_json"])} for r in rows]


def sample_rows(db, batch_id, limit=SAMPLE_ROWS):
    return db.execute(
        "SELECT * FROM medalist_import_staging WHERE batch_id = ?"
        " AND status = 'valid' ORDER BY row_no LIMIT ?",
        (batch_id, limit)).fetchall()


def list_batches(db, convocation_id, limit=20):
    return db.execute(
        "SELECT * FROM medalist_import_batches WHERE convocation_id = ?"
        " ORDER BY id DESC LIMIT ?", (convocation_id, limit)).fetchall()
